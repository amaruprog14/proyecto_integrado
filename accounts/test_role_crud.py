"""
Pruebas del CRUD de Roles.

Cubren creación, consulta, edición, desactivación y reactivación lógicas, validaciones,
permisos (403 con auditoría), búsqueda, filtros, paginación y exportación a Excel, además de
la coherencia con el CRUD de usuarios: roles inactivos fuera del formulario de usuarios,
advertencia por funcionarios afectados, protección del último Administrador global y
auditoría de cambios de permisos y alcance.
"""

import re
from io import BytesIO
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.forms import OfficialForm, OfficialUserForm
from accounts.models import Delegation, Official, Role
from accounts.role_forms import last_global_administrator
from configuration.models import AuditLog

User = get_user_model()
PASSWORD = "Clave-Segura-123"


def permission(codename):
    app_label, name = codename.split(".")
    return Permission.objects.get(content_type__app_label=app_label, codename=name)


def grant_permissions(user, *codenames):
    for codename in codenames:
        user.user_permissions.add(permission(codename))
    return user


def make_official(user, *, delegation=None, status=Official.Status.ACTIVE, full_name=None, national_id=None, email=None):
    return Official.objects.create(
        user=user,
        full_name=full_name or user.username.replace(".", " ").title(),
        national_id=national_id or f"RUT-{user.username}",
        email=email or f"{user.username}@example.test",
        delegation=delegation,
        status=status,
    )


def make_role(name, *codenames, scope=Role.Scope.DELEGATION, is_active=True, description=""):
    role = Role.objects.create(role_name=name, scope=scope, is_active=is_active, description=description)
    role.permissions.set([permission(codename) for codename in codenames])
    return role


def role_payload(*codenames, **overrides):
    payload = {
        "role_name": "Rol de Prueba",
        "description": "Rol creado en las pruebas",
        "scope": Role.Scope.DELEGATION,
        "permissions": [permission(codename).pk for codename in codenames or ("accounts.view_official",)],
        "is_active": "True",
    }
    payload.update(overrides)
    return payload


def official_with_role(username, role, *, status=Official.Status.ACTIVE, delegation=None):
    user = User.objects.create_user(username, f"{username}@example.test", PASSWORD)
    official = make_official(user, status=status, delegation=delegation, national_id=username)
    official.roles.add(role)
    return official


class RoleCrudAdminTests(TestCase):
    """CRUD de roles con un superusuario."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("admin.roles", "admin.roles@example.test", PASSWORD)
        make_official(cls.admin, full_name="Administrador Roles", national_id="RUT-ADMIN-ROLES")
        cls.role = make_role(
            "Operador Territorial",
            "activities.view_activity",
            "activities.add_activity",
            description="Registra actividades en terreno",
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_list_renders_persisted_roles(self):
        response = self.client.get(reverse("role_list"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "accounts/role_list.html")
        self.assertContains(response, "Operador Territorial")
        self.assertContains(response, "Delegaciones autorizadas")
        self.assertContains(response, f'href="{reverse("role_create")}"')
        self.assertEqual(response.context["roles"][0].permission_count, 2)

    def test_create_persists_scope_permissions_and_status(self):
        response = self.client.post(
            reverse("role_create"),
            role_payload(
                "configuration.view_item",
                "configuration.change_item",
                role_name="Gestor de Ítems",
                scope=Role.Scope.ALL,
                is_active="False",
            ),
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        role = Role.objects.get(role_name="Gestor de Ítems")
        self.assertEqual(role.scope, Role.Scope.ALL)
        self.assertFalse(role.is_active)
        self.assertEqual(
            set(role.permissions.values_list("codename", flat=True)), {"view_item", "change_item"}
        )
        self.assertContains(response, "fue creado correctamente")
        log = AuditLog.objects.get(action="Creación de rol", affected_record_id=role.pk)
        self.assertIn("configuration.change_item", log.new_value)

    def test_name_extra_spaces_are_removed(self):
        self.client.post(reverse("role_create"), role_payload(role_name="  Analista    de   Datos  "))

        self.assertTrue(Role.objects.filter(role_name="Analista de Datos").exists())

    def test_duplicate_name_is_rejected_case_insensitively(self):
        response = self.client.post(
            reverse("role_create"), role_payload(role_name="  operador   TERRITORIAL ")
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ya existe un rol con este nombre")
        self.assertEqual(Role.objects.filter(role_name__iexact="Operador Territorial").count(), 1)

    def test_required_fields_are_validated(self):
        response = self.client.post(
            reverse("role_create"),
            {"role_name": "   ", "description": "", "scope": "", "is_active": "True"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ingresa el nombre del rol")
        self.assertContains(response, "Selecciona el alcance del rol")
        self.assertContains(response, "Selecciona al menos un permiso")
        self.assertEqual(Role.objects.count(), 1)

    def test_delete_permissions_are_not_offered(self):
        response = self.client.post(
            reverse("role_create"),
            role_payload("activities.delete_activity", role_name="Rol con Borrado"),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no está disponible")
        self.assertFalse(Role.objects.filter(role_name="Rol con Borrado").exists())

    def test_update_persists_changes(self):
        response = self.client.post(
            reverse("role_update", args=[self.role.pk]),
            role_payload(
                "activities.view_activity",
                role_name="Operador Territorial Editado",
                description="Nueva descripción",
                scope=Role.Scope.OWN,
            ),
            follow=True,
        )

        self.role.refresh_from_db()
        self.assertEqual(self.role.role_name, "Operador Territorial Editado")
        self.assertEqual(self.role.description, "Nueva descripción")
        self.assertEqual(self.role.scope, Role.Scope.OWN)
        self.assertEqual(list(self.role.permissions.values_list("codename", flat=True)), ["view_activity"])
        self.assertContains(response, "fue actualizado correctamente")

    def test_update_audits_previous_and_new_permissions_and_scope(self):
        self.client.post(
            reverse("role_update", args=[self.role.pk]),
            role_payload(
                "activities.view_activity",
                "configuration.view_item",
                role_name=self.role.role_name,
                description=self.role.description,
                scope=Role.Scope.ALL,
            ),
        )

        log = AuditLog.objects.get(action="Actualización de rol", affected_record_id=self.role.pk)
        self.assertEqual(log.affected_table, Role._meta.db_table)
        self.assertIn("Alcance: Delegaciones autorizadas", log.previous_value)
        self.assertIn("Alcance: Transversal (todas las delegaciones)", log.new_value)
        self.assertIn("activities.add_activity", log.previous_value)
        self.assertNotIn("activities.add_activity", log.new_value)
        self.assertIn("configuration.view_item", log.new_value)
        self.assertNotIn("Nombre:", log.new_value)

    def test_update_keeps_permissions_outside_the_catalog(self):
        self.role.permissions.add(permission("activities.delete_activity"))

        self.client.post(
            reverse("role_update", args=[self.role.pk]),
            role_payload("activities.view_activity", role_name=self.role.role_name),
        )

        self.assertEqual(
            set(self.role.permissions.values_list("codename", flat=True)),
            {"view_activity", "delete_activity"},
        )

    def test_update_form_does_not_change_status(self):
        self.client.post(
            reverse("role_update", args=[self.role.pk]),
            role_payload("activities.view_activity", role_name=self.role.role_name, is_active="False"),
        )

        self.role.refresh_from_db()
        self.assertTrue(self.role.is_active)

    def test_deactivate_role_with_active_officials_warns_and_audits(self):
        official = official_with_role("operador.uno", self.role)
        official_with_role("operador.dos", self.role)
        official_with_role("operador.inactivo", self.role, status=Official.Status.INACTIVE)
        self.assertTrue(User.objects.get(pk=official.user_id).has_perm("activities.add_activity"))

        response = self.client.post(
            reverse("role_deactivate", args=[self.role.pk]), {"affected_officials": "2"}, follow=True
        )

        self.role.refresh_from_db()
        self.assertFalse(self.role.is_active)
        self.assertTrue(Role.objects.filter(pk=self.role.pk).exists())
        self.assertContains(response, "2 funcionarios activos quedaron sin los permisos de este rol")
        log = AuditLog.objects.get(action="Desactivación de rol", affected_record_id=self.role.pk)
        self.assertIn("pierden los permisos del rol: 2", log.new_value)
        self.assertFalse(User.objects.get(pk=official.user_id).has_perm("activities.add_activity"))

    def test_deactivate_confirmation_shows_affected_officials(self):
        official_with_role("operador.uno", self.role)

        response = self.client.get(reverse("role_list"))

        self.assertContains(response, "1 funcionario activo quedará sin los permisos de este rol")
        self.assertContains(response, 'name="affected_officials" value="1"')

    def test_deactivate_with_outdated_count_is_not_applied(self):
        official_with_role("operador.uno", self.role)

        response = self.client.post(
            reverse("role_deactivate", args=[self.role.pk]), {"affected_officials": "0"}, follow=True
        )

        self.role.refresh_from_db()
        self.assertTrue(self.role.is_active)
        self.assertContains(response, "confirma nuevamente")
        self.assertFalse(AuditLog.objects.filter(action="Desactivación de rol").exists())

    def test_deactivate_rejects_get_requests(self):
        response = self.client.get(reverse("role_deactivate", args=[self.role.pk]))

        self.assertEqual(response.status_code, 405)
        self.role.refresh_from_db()
        self.assertTrue(self.role.is_active)

    def test_activate_after_deactivate(self):
        self.role.is_active = False
        self.role.save()

        response = self.client.post(reverse("role_activate", args=[self.role.pk]), follow=True)

        self.role.refresh_from_db()
        self.assertTrue(self.role.is_active)
        self.assertContains(response, "fue activado correctamente")
        self.assertTrue(AuditLog.objects.filter(action="Activación de rol", affected_record_id=self.role.pk).exists())

    def test_search_is_partial_and_case_insensitive_across_fields(self):
        make_role("Supervisor Comunal", "accounts.view_delegation", scope=Role.Scope.ALL)

        cases = (
            ("TERRITORIAL", "Operador Territorial", "Supervisor Comunal"),
            ("en terreno", "Operador Territorial", "Supervisor Comunal"),
            ("transversal", "Supervisor Comunal", "Operador Territorial"),
            ("view_delegation", "Supervisor Comunal", "Operador Territorial"),
        )
        for query, expected, excluded in cases:
            with self.subTest(query=query):
                response = self.client.get(reverse("role_list"), {"q": query})
                self.assertContains(response, expected)
                self.assertNotContains(response, excluded)

    def test_status_filter_combines_with_search(self):
        make_role("Operador Rural", "activities.view_activity", is_active=False)

        response = self.client.get(reverse("role_list"), {"q": "operador", "status": "INACTIVE"})

        self.assertContains(response, "Operador Rural")
        self.assertNotContains(response, "Operador Territorial")

    def test_empty_states_distinguish_no_matches(self):
        response = self.client.get(reverse("role_list"), {"q": "no-existe"})

        self.assertContains(response, "No hay roles que coincidan con la búsqueda o el filtro aplicado.")

    def test_pagination_is_ten_per_page_and_keeps_filters(self):
        for index in range(12):
            make_role(f"Rol Paginado {index:02d}", "activities.view_activity")

        first_page = self.client.get(reverse("role_list"), {"q": "paginado"})

        self.assertEqual(len(first_page.context["roles"]), 10)
        self.assertEqual(first_page.context["paginator"].count, 12)
        self.assertEqual(first_page.context["querystring"], "q=paginado")
        self.assertContains(first_page, "q=paginado&amp;page=2")

        second_page = self.client.get(reverse("role_list"), {"q": "paginado", "page": 2})

        self.assertEqual(len(second_page.context["roles"]), 2)

    def test_out_of_range_page_is_clamped(self):
        response = self.client.get(reverse("role_list"), {"page": 999})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_obj"].number, 1)

    def test_export_is_xlsx_with_filtered_rows_and_no_internal_fields(self):
        for index in range(12):
            make_role(f"Rol Exportado {index:02d}", "activities.view_activity")
        make_role("Rol Exportado Inactivo", "activities.view_activity", is_active=False)

        response = self.client.get(reverse("role_export"), {"q": "exportado", "status": "ACTIVE"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('attachment; filename="roles.xlsx"', response["Content-Disposition"])
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(
            rows[0], ("Nombre", "Descripción", "Alcance", "Permisos", "Funcionarios activos", "Estado")
        )
        self.assertEqual(len(rows), 13)
        self.assertNotIn("Rol Exportado Inactivo", [row[0] for row in rows])
        self.assertEqual(rows[1][3], "Actividades: Consultar")
        self.assertEqual(rows[1][5], "Activo")

    def test_export_stores_formula_like_values_as_text(self):
        make_role("=1+1", "activities.view_activity")

        response = self.client.get(reverse("role_export"), {"q": "=1+1"})

        cell = load_workbook(BytesIO(response.content)).active["A2"]
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_export_file_contains_no_formula_for_names_starting_with_equals(self):
        make_role('=HYPERLINK("http://ejemplo.test","clic")', "activities.view_activity")

        response = self.client.get(reverse("role_export"), {"q": "=HYPERLINK"})

        with ZipFile(BytesIO(response.content)) as workbook:
            sheet_xml = workbook.read("xl/worksheets/sheet1.xml")
        self.assertIsNone(re.search(rb"<f[ >]", sheet_xml))
        self.assertIn(b'<c r="A2" t="inlineStr">', sheet_xml)
        self.assertIn(b"=HYPERLINK", sheet_xml)

    def test_create_update_and_detail_views_render(self):
        official_with_role("operador.detalle", self.role)

        create_response = self.client.get(reverse("role_create"))
        self.assertEqual(create_response.status_code, 200)
        self.assertTemplateUsed(create_response, "accounts/role_form.html")

        update_response = self.client.get(reverse("role_update", args=[self.role.pk]))
        self.assertEqual(update_response.status_code, 200)
        self.assertContains(update_response, "Operador Territorial")
        self.assertContains(update_response, "Actividades")

        detail_response = self.client.get(reverse("role_detail", args=[self.role.pk]))
        self.assertEqual(detail_response.status_code, 200)
        self.assertTemplateUsed(detail_response, "accounts/role_detail.html")
        self.assertContains(detail_response, "Consultar, Crear")
        self.assertContains(detail_response, "Operador Detalle")

    def test_anonymous_is_redirected_to_login(self):
        self.client.logout()
        url = reverse("role_list")

        self.assertRedirects(self.client.get(url), f"{reverse('login')}?next={url}")


class GlobalAdministratorProtectionTests(TestCase):
    """El sistema nunca queda sin un rol Administrador global con funcionarios activos."""

    @classmethod
    def setUpTestData(cls):
        cls.superuser = User.objects.create_superuser("super.roles", "super.roles@example.test", PASSWORD)
        cls.admin_role = make_role(
            "Administrador Global",
            "accounts.view_role",
            "accounts.change_role",
            "auth.change_user",
            scope=Role.Scope.ALL,
        )
        cls.admin_official = official_with_role("admin.global", cls.admin_role)

    def setUp(self):
        self.client.force_login(self.superuser)

    def edit_payload(self, *codenames, **overrides):
        payload = role_payload(
            *(codenames or ("accounts.view_role", "accounts.change_role", "auth.change_user")),
            role_name=self.admin_role.role_name,
            scope=Role.Scope.ALL,
        )
        payload.update(overrides)
        return payload

    def test_last_global_administrator_cannot_be_deactivated(self):
        response = self.client.post(
            reverse("role_deactivate", args=[self.admin_role.pk]),
            {"affected_officials": "1"},
            follow=True,
        )

        self.admin_role.refresh_from_db()
        self.assertTrue(self.admin_role.is_active)
        self.assertContains(response, "es el último rol Administrador global con funcionarios activos")

    def test_last_global_administrator_cannot_lose_privileged_permissions(self):
        response = self.client.post(
            reverse("role_update", args=[self.admin_role.pk]), self.edit_payload("accounts.view_role")
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "quitar los permisos privilegiados ni cambiar el alcance")
        self.assertTrue(self.admin_role.permissions.filter(codename="change_role").exists())

    def test_last_global_administrator_cannot_change_scope(self):
        response = self.client.post(
            reverse("role_update", args=[self.admin_role.pk]),
            self.edit_payload(scope=Role.Scope.DELEGATION),
        )

        self.assertEqual(response.status_code, 200)
        self.admin_role.refresh_from_db()
        self.assertEqual(self.admin_role.scope, Role.Scope.ALL)

    def test_last_global_administrator_can_still_be_renamed(self):
        self.client.post(
            reverse("role_update", args=[self.admin_role.pk]), self.edit_payload(role_name="Administración")
        )

        self.admin_role.refresh_from_db()
        self.assertEqual(self.admin_role.role_name, "Administración")

    def test_list_disables_deactivation_of_last_global_administrator(self):
        response = self.client.get(reverse("role_list"))

        self.assertNotContains(response, reverse("role_deactivate", args=[self.admin_role.pk]))

    def test_deactivation_is_allowed_when_another_global_administrator_is_active(self):
        backup = make_role("Administrador Respaldo", "accounts.change_role", scope=Role.Scope.ALL)
        official_with_role("admin.respaldo", backup)

        self.client.post(reverse("role_deactivate", args=[self.admin_role.pk]), {"affected_officials": "1"})

        self.admin_role.refresh_from_db()
        self.assertFalse(self.admin_role.is_active)

    def test_backup_without_active_officials_does_not_count(self):
        backup = make_role("Administrador Sin Personal", "accounts.change_role", scope=Role.Scope.ALL)
        official_with_role("admin.suspendido", backup, status=Official.Status.SUSPENDED)

        self.client.post(reverse("role_deactivate", args=[self.admin_role.pk]), {"affected_officials": "1"})

        self.admin_role.refresh_from_db()
        self.assertTrue(self.admin_role.is_active)

    def test_last_global_administrator_is_none_with_two_staffed_global_administrators(self):
        official_with_role("admin.global.dos", self.admin_role)
        self.assertEqual(last_global_administrator(), self.admin_role)

        backup = make_role("Administrador Respaldo", "accounts.change_role", scope=Role.Scope.ALL)
        official_with_role("admin.respaldo", backup)

        self.assertIsNone(last_global_administrator())

    def test_role_cannot_become_global_administrator_while_combined_with_other_roles(self):
        operator = make_role("Operador", "activities.view_activity")
        auditor = make_role("Auditor", "configuration.view_auditlog")
        official = official_with_role("combinado", operator)
        official.roles.add(auditor)

        response = self.client.post(
            reverse("role_update", args=[operator.pk]),
            role_payload("activities.view_activity", "auth.change_user", role_name="Operador", scope=Role.Scope.ALL),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no puede combinarse con otros roles")
        self.assertFalse(operator.permissions.filter(codename="change_user").exists())

    def test_global_administrator_cannot_be_reactivated_while_combined_with_other_roles(self):
        backup = make_role("Administrador Inactivo", "accounts.change_role", scope=Role.Scope.ALL, is_active=False)
        official = official_with_role("combinado.inactivo", backup)
        official.roles.add(make_role("Consulta", "activities.view_activity"))

        response = self.client.post(reverse("role_activate", args=[backup.pk]), follow=True)

        backup.refresh_from_db()
        self.assertFalse(backup.is_active)
        self.assertContains(response, "no puede combinarse con otros roles")


class RoleUserFormIntegrationTests(TestCase):
    """Los formularios de usuarios y funcionarios reflejan el estado de los roles."""

    @classmethod
    def setUpTestData(cls):
        cls.superuser = User.objects.create_superuser("super.forms", "super.forms@example.test", PASSWORD)
        make_official(cls.superuser, national_id="RUT-SUPER-FORMS")
        cls.role = make_role("Verificador de Evidencias", "activities.view_activity")

    def setUp(self):
        self.client.force_login(self.superuser)

    def test_deactivated_role_disappears_from_user_forms(self):
        self.assertIn(self.role, OfficialUserForm(actor=self.superuser).fields["roles"].queryset)
        self.assertContains(self.client.get(reverse("user_create")), "Verificador de Evidencias")

        self.client.post(
            reverse("role_deactivate", args=[self.role.pk]), {"affected_officials": "0"}, follow=True
        )

        self.assertNotIn(self.role, OfficialUserForm(actor=self.superuser).fields["roles"].queryset)
        self.assertNotIn(self.role, OfficialForm(actor=self.superuser).fields["roles"].queryset)
        self.assertNotContains(self.client.get(reverse("user_create")), "Verificador de Evidencias")

        self.client.post(reverse("role_activate", args=[self.role.pk]))

        self.assertIn(self.role, OfficialUserForm(actor=self.superuser).fields["roles"].queryset)

    def test_role_saved_as_global_administrator_is_flagged_in_user_form(self):
        self.client.post(
            reverse("role_create"),
            role_payload("accounts.change_delegation", role_name="Administración Central", scope=Role.Scope.ALL),
        )
        role = Role.objects.get(role_name="Administración Central")

        response = self.client.get(reverse("user_create"))

        self.assertContains(response, f'value="{role.pk}"')
        self.assertContains(response, 'data-role-global-administrator="true"')


class RolePermissionTests(TestCase):
    """Cada operación exige su permiso; los rechazos responden 403 y quedan auditados."""

    @classmethod
    def setUpTestData(cls):
        cls.role = make_role("Rol Existente", "activities.view_activity")

        cls.viewer = User.objects.create_user("viewer.roles", "viewer.roles@example.test", PASSWORD)
        make_official(cls.viewer)
        grant_permissions(cls.viewer, "accounts.view_role")

        cls.creator = User.objects.create_user("creator.roles", "creator.roles@example.test", PASSWORD)
        make_official(cls.creator)
        grant_permissions(cls.creator, "accounts.view_role", "accounts.add_role")

        cls.stranger = User.objects.create_user("stranger.roles", "stranger.roles@example.test", PASSWORD)
        make_official(cls.stranger)

    def test_viewer_can_list_and_export_but_not_modify(self):
        self.client.force_login(self.viewer)

        self.assertEqual(self.client.get(reverse("role_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("role_export")).status_code, 200)
        self.assertEqual(self.client.post(reverse("role_create"), role_payload()).status_code, 403)
        self.assertEqual(self.client.get(reverse("role_update", args=[self.role.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse("role_deactivate", args=[self.role.pk])).status_code, 403)
        self.assertNotContains(self.client.get(reverse("role_list")), reverse("role_create"))

        self.role.refresh_from_db()
        self.assertTrue(self.role.is_active)
        self.assertFalse(Role.objects.filter(role_name="Rol de Prueba").exists())
        self.assertEqual(AuditLog.objects.filter(action="Acceso denegado").count(), 3)

    def test_creator_without_change_role_can_create_regular_roles(self):
        self.client.force_login(self.creator)

        self.client.post(reverse("role_create"), role_payload("activities.view_activity"))

        self.assertTrue(Role.objects.filter(role_name="Rol de Prueba").exists())

    def test_creator_without_change_role_cannot_grant_privileged_permissions(self):
        self.client.force_login(self.creator)

        response = self.client.post(
            reverse("role_create"), role_payload("auth.change_user", role_name="Rol Escalado")
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Role.objects.filter(role_name="Rol Escalado").exists())
        log = AuditLog.objects.get(action="Acceso denegado")
        self.assertIn("accounts.change_role", log.new_value)
        self.assertEqual(log.official, self.creator.official)

    def test_privileged_permissions_are_denied_even_if_the_submission_has_other_errors(self):
        self.client.force_login(self.creator)

        response = self.client.post(
            reverse("role_create"), role_payload("auth.change_user", role_name="")
        )

        self.assertEqual(response.status_code, 403)
        self.assertTrue(AuditLog.objects.filter(action="Acceso denegado").exists())

    def test_privileged_permissions_are_locked_in_the_creator_form(self):
        self.client.force_login(self.creator)

        response = self.client.get(reverse("role_create"))

        self.assertContains(response, 'data-privileged="true" disabled')

    def test_user_without_permission_is_rejected_and_audited(self):
        self.client.force_login(self.stranger)

        self.assertEqual(self.client.get(reverse("role_list")).status_code, 403)
        self.assertTrue(AuditLog.objects.filter(action="Acceso denegado").exists())

    def test_deactivating_a_role_removes_the_access_it_granted(self):
        manager_role = make_role("Gestor de Roles", "accounts.view_role")
        official = official_with_role("gestor.roles", manager_role)
        self.client.force_login(official.user)
        self.assertEqual(self.client.get(reverse("role_list")).status_code, 200)

        manager_role.is_active = False
        manager_role.save()

        self.assertEqual(self.client.get(reverse("role_list")).status_code, 403)


class RoleDetailScopeTests(TestCase):
    """El detalle solo lista funcionarios de las delegaciones del ámbito del actor."""

    @classmethod
    def setUpTestData(cls):
        cls.own = Delegation.objects.create(delegation_name="Zona Propia", address="Calle 1", commune="La Serena")
        cls.foreign = Delegation.objects.create(delegation_name="Zona Ajena", address="Calle 2", commune="La Serena")
        cls.role = make_role("Operador Compartido", "activities.view_activity")
        official_with_role("operador.propio", cls.role, delegation=cls.own)
        official_with_role("operador.ajeno", cls.role, delegation=cls.foreign)

        cls.viewer = User.objects.create_user("viewer.scope.roles", "viewer.scope.roles@example.test", PASSWORD)
        make_official(cls.viewer, delegation=cls.own, national_id="RUT-VIEWER-SCOPE")
        grant_permissions(cls.viewer, "accounts.view_role")

    def test_detail_lists_only_officials_in_scope(self):
        self.client.force_login(self.viewer)

        response = self.client.get(reverse("role_detail", args=[self.role.pk]))

        self.assertContains(response, "Operador Propio")
        self.assertNotContains(response, "Operador Ajeno")
        self.assertEqual(response.context["active_official_count"], 2)
