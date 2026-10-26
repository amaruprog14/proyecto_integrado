"""
Pruebas del CRUD de Delegaciones.

Cubren creación, consulta, edición, desactivación lógica, validaciones, permisos,
búsqueda, filtros combinables, paginación, exportación Excel y el ámbito por delegación
(SEG-01/SEG-02: acceso directo a identificadores fuera del ámbito).
"""

from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Delegation, Official, Role
from configuration.models import AuditLog

User = get_user_model()
PASSWORD = "Clave-Segura-123"


def grant_permissions(user, *codenames):
    for codename in codenames:
        app_label, name = codename.split(".")
        permission = Permission.objects.get(content_type__app_label=app_label, codename=name)
        user.user_permissions.add(permission)
    return user


def make_official(user, *, delegation=None, status=Official.Status.ACTIVE, full_name=None, national_id=None, email=None):
    return Official.objects.create(
        user=user,
        full_name=full_name or user.username.replace(".", " ").title(),
        national_id=national_id or f"RUT-{user.username}",
        email=email or f"{user.username}@example.test",
        delegation=delegation,
        status=status,
    )


def delegation_payload(**overrides):
    payload = {
        "delegation_name": "Zona de Prueba",
        "address": "Calle de Prueba 100",
        "commune": "La Serena",
        "status": Delegation.Status.ACTIVE,
    }
    payload.update(overrides)
    return payload


class DelegationCrudAdminTests(TestCase):
    """CRUD de delegaciones con un administrador de alcance total."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("admin.crud", "admin.crud@example.test", PASSWORD)
        make_official(cls.admin, full_name="Administrador CRUD", national_id="RUT-ADMIN-CRUD")
        cls.delegation = Delegation.objects.create(
            delegation_name="Zona Rural Norte",
            address="Camino Rural 10",
            commune="La Serena",
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_list_renders_persisted_records(self):
        response = self.client.get(reverse("delegation_list"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "accounts/delegation_list.html")
        self.assertContains(response, "Zona Rural Norte")
        self.assertContains(response, self.delegation.code)
        self.assertContains(response, f'href="{reverse("delegation_create")}"')

    def test_create_persists_and_generates_code(self):
        response = self.client.post(
            reverse("delegation_create"),
            delegation_payload(delegation_name="Zona Nueva"),
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        delegation = Delegation.objects.get(delegation_name="Zona Nueva")
        self.assertEqual(delegation.code, f"DEL-{delegation.pk:04d}")
        self.assertContains(response, "fue creada correctamente")

    def test_duplicate_name_is_rejected_case_insensitively(self):
        response = self.client.post(
            reverse("delegation_create"),
            delegation_payload(delegation_name="zona rural norte"),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ya existe una delegación con este nombre")
        self.assertEqual(Delegation.objects.filter(delegation_name__iexact="Zona Rural Norte").count(), 1)

    def test_required_fields_are_validated(self):
        response = self.client.post(
            reverse("delegation_create"),
            {"delegation_name": "   ", "address": "", "commune": "", "status": Delegation.Status.ACTIVE},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ingresa el nombre de la delegación")
        self.assertContains(response, "Ingresa la dirección")
        self.assertContains(response, "Ingresa la comuna")
        self.assertEqual(Delegation.objects.count(), 1)

    def test_update_persists_changes_and_keeps_code(self):
        original_code = self.delegation.code
        response = self.client.post(
            reverse("delegation_update", args=[self.delegation.pk]),
            delegation_payload(
                delegation_name="Zona Rural Norte Editada",
                address="Camino Rural 20",
                status=Delegation.Status.INACTIVE,
            ),
            follow=True,
        )

        self.delegation.refresh_from_db()
        self.assertEqual(self.delegation.delegation_name, "Zona Rural Norte Editada")
        self.assertEqual(self.delegation.address, "Camino Rural 20")
        self.assertEqual(self.delegation.status, Delegation.Status.INACTIVE)
        self.assertEqual(self.delegation.code, original_code)
        self.assertContains(response, "fue actualizada correctamente")

    def test_deactivate_is_logical_and_keeps_history(self):
        response = self.client.post(
            reverse("delegation_deactivate", args=[self.delegation.pk]), follow=True
        )

        self.delegation.refresh_from_db()
        self.assertEqual(self.delegation.status, Delegation.Status.INACTIVE)
        self.assertTrue(Delegation.objects.filter(pk=self.delegation.pk).exists())
        self.assertContains(response, "fue desactivada correctamente")
        self.assertTrue(AuditLog.objects.filter(action="Desactivación de delegación").exists())

    def test_deactivate_does_not_remove_delegation_even_with_relations(self):
        other = User.objects.create_user("relation.user", "relation.user@example.test", PASSWORD)
        make_official(other, delegation=self.delegation, national_id="RUT-RELATION")

        self.client.post(reverse("delegation_deactivate", args=[self.delegation.pk]))

        self.delegation.refresh_from_db()
        self.assertEqual(self.delegation.status, Delegation.Status.INACTIVE)
        self.assertTrue(Delegation.objects.filter(pk=self.delegation.pk).exists())

    def test_activate_after_deactivate(self):
        self.delegation.deactivate()

        response = self.client.post(
            reverse("delegation_activate", args=[self.delegation.pk]), follow=True
        )

        self.delegation.refresh_from_db()
        self.assertEqual(self.delegation.status, Delegation.Status.ACTIVE)
        self.assertContains(response, "fue activada correctamente")

    def test_deactivate_rejects_get_requests(self):
        response = self.client.get(reverse("delegation_deactivate", args=[self.delegation.pk]))

        self.assertEqual(response.status_code, 405)
        self.delegation.refresh_from_db()
        self.assertEqual(self.delegation.status, Delegation.Status.ACTIVE)

    def test_search_is_partial_and_case_insensitive(self):
        Delegation.objects.create(
            delegation_name="Zona Urbana Sur", address="Calle Sur 500", commune="La Serena"
        )

        response = self.client.get(reverse("delegation_list"), {"q": "rural norte"})

        names = [item.delegation_name for item in response.context["delegations"]]
        self.assertIn("Zona Rural Norte", names)
        self.assertNotIn("Zona Urbana Sur", names)

    def test_status_filter_combines_with_search(self):
        Delegation.objects.create(
            delegation_name="Zona Rural Sur",
            address="Camino Sur 5",
            commune="La Serena",
            status=Delegation.Status.INACTIVE,
        )

        response = self.client.get(
            reverse("delegation_list"), {"q": "rural", "status": Delegation.Status.INACTIVE}
        )

        names = [item.delegation_name for item in response.context["delegations"]]
        self.assertIn("Zona Rural Sur", names)
        self.assertNotIn("Zona Rural Norte", names)

    def test_pagination_is_ten_per_page_and_keeps_filters(self):
        for index in range(12):
            Delegation.objects.create(
                delegation_name=f"Delegación {index:02d}",
                address="Calle Genérica",
                commune="La Serena",
            )

        first_page = self.client.get(reverse("delegation_list"), {"status": Delegation.Status.ACTIVE})

        self.assertEqual(len(first_page.context["delegations"]), 10)
        self.assertEqual(first_page.context["paginator"].count, 13)
        self.assertEqual(first_page.context["querystring"], "status=ACTIVE")

        second_page = self.client.get(
            reverse("delegation_list"), {"status": Delegation.Status.ACTIVE, "page": 2}
        )

        self.assertEqual(len(second_page.context["delegations"]), 3)
        self.assertEqual(second_page.context["querystring"], "status=ACTIVE")

    def test_out_of_range_page_is_clamped(self):
        response = self.client.get(reverse("delegation_list"), {"page": 999})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_obj"].number, 1)

    def test_export_respects_filters_and_is_not_paginated(self):
        for index in range(12):
            Delegation.objects.create(
                delegation_name=f"Delegación {index:02d}",
                address="Calle Genérica",
                commune="La Serena",
            )

        response = self.client.get(reverse("delegation_export"), {"status": Delegation.Status.ACTIVE})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('attachment; filename="delegaciones.xlsx"', response["Content-Disposition"])

        sheet = load_workbook(BytesIO(response.content), read_only=True).active
        rows = list(sheet.values)
        self.assertEqual(len(rows), 14)
        self.assertEqual(rows[0], ("Código", "Nombre", "Dirección", "Comuna", "Estado"))
        self.assertNotIn("password", str(rows).lower())

    def test_export_rows_follow_search_and_status_without_internal_fields(self):
        expected = Delegation.objects.create(
            delegation_name="Zona Costera Sur", address="Avenida del Mar 200", commune="Coquimbo"
        )
        Delegation.objects.create(
            delegation_name="Zona Costera Norte",
            address="Avenida del Mar 300",
            commune="La Serena",
            status=Delegation.Status.INACTIVE,
        )

        response = self.client.get(
            reverse("delegation_export"), {"q": "COSTERA", "status": Delegation.Status.ACTIVE}
        )

        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        # Solo las columnas visibles: sin id interno, fechas de auditoría ni borrado lógico.
        self.assertEqual(
            rows,
            [
                ("Código", "Nombre", "Dirección", "Comuna", "Estado"),
                (expected.code, "Zona Costera Sur", "Avenida del Mar 200", "Coquimbo", "Activa"),
            ],
        )

    def test_export_contains_every_filtered_record_shown_across_pages(self):
        for index in range(23):
            Delegation.objects.create(
                delegation_name=f"Sector Exportable {index:02d}", address="Calle Larga", commune="La Serena"
            )

        listed = self.client.get(reverse("delegation_list"), {"q": "exportable"})
        export = self.client.get(reverse("delegation_export"), {"q": "exportable"})

        self.assertEqual(listed.context["total"], 23)
        self.assertEqual(listed.context["paginator"].num_pages, 3)
        names = [row[1] for row in load_workbook(BytesIO(export.content), read_only=True).active.values][1:]
        self.assertEqual(names, [f"Sector Exportable {index:02d}" for index in range(23)])

    def test_anonymous_is_redirected_to_login(self):
        self.client.logout()
        url = reverse("delegation_list")

        self.assertRedirects(self.client.get(url), f"{reverse('login')}?next={url}")

    def test_create_update_and_detail_views_render(self):
        create_response = self.client.get(
            reverse("delegation_create"), HTTP_X_MODAL_REQUEST="true"
        )
        self.assertEqual(create_response.status_code, 200)
        self.assertTemplateUsed(create_response, "accounts/delegation_form.html")

        update_response = self.client.get(
            reverse("delegation_update", args=[self.delegation.pk]), HTTP_X_MODAL_REQUEST="true"
        )
        self.assertEqual(update_response.status_code, 200)
        self.assertTemplateUsed(update_response, "accounts/delegation_form.html")
        self.assertContains(update_response, self.delegation.code)

        detail_response = self.client.get(reverse("delegation_detail", args=[self.delegation.pk]))
        self.assertEqual(detail_response.status_code, 200)
        self.assertTemplateUsed(detail_response, "accounts/delegation_detail.html")
        self.assertContains(detail_response, "Zona Rural Norte")
        self.assertContains(detail_response, "Camino Rural 10")


class DelegationScopeTests(TestCase):
    """Un funcionario solo ve y administra las delegaciones de su ámbito."""

    @classmethod
    def setUpTestData(cls):
        cls.authorized = Delegation.objects.create(
            delegation_name="Zona Autorizada", address="Calle Uno", commune="La Serena"
        )
        cls.foreign = Delegation.objects.create(
            delegation_name="Zona Ajena", address="Calle Dos", commune="La Serena"
        )

        cls.supervisor = User.objects.create_user(
            "supervisor.scope", "supervisor.scope@example.test", PASSWORD
        )
        make_official(
            cls.supervisor,
            delegation=cls.authorized,
            national_id="RUT-SUPERVISOR",
            email="supervisor.scope@example.test",
        )
        grant_permissions(
            cls.supervisor,
            "accounts.view_delegation",
            "accounts.add_delegation",
            "accounts.change_delegation",
        )

    def setUp(self):
        self.client.force_login(self.supervisor)

    def test_list_only_shows_authorized_delegations(self):
        response = self.client.get(reverse("delegation_list"))

        self.assertContains(response, "Zona Autorizada")
        self.assertNotContains(response, "Zona Ajena")

    def test_foreign_delegation_detail_is_not_visible(self):
        response = self.client.get(reverse("delegation_detail", args=[self.foreign.pk]))

        self.assertEqual(response.status_code, 404)

    def test_foreign_delegation_cannot_be_updated(self):
        response = self.client.post(
            reverse("delegation_update", args=[self.foreign.pk]),
            delegation_payload(delegation_name="Zona Secuestrada"),
        )

        self.assertEqual(response.status_code, 404)
        self.foreign.refresh_from_db()
        self.assertEqual(self.foreign.delegation_name, "Zona Ajena")

    def test_foreign_delegation_cannot_be_deactivated(self):
        response = self.client.post(reverse("delegation_deactivate", args=[self.foreign.pk]))

        self.assertEqual(response.status_code, 404)
        self.foreign.refresh_from_db()
        self.assertEqual(self.foreign.status, Delegation.Status.ACTIVE)

    def test_export_only_includes_authorized_delegations(self):
        response = self.client.get(reverse("delegation_export"))
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        content = str(rows)

        self.assertIn("Zona Autorizada", content)
        self.assertNotIn("Zona Ajena", content)


class DelegationPermissionTests(TestCase):
    """Permisos estándar de Django controlan cada operación en el servidor."""

    @classmethod
    def setUpTestData(cls):
        cls.delegation = Delegation.objects.create(
            delegation_name="Zona Permisos", address="Calle Uno", commune="La Serena"
        )

        cls.viewer = User.objects.create_user("viewer.scope", "viewer.scope@example.test", PASSWORD)
        make_official(cls.viewer, delegation=cls.delegation, national_id="RUT-VIEWER")
        grant_permissions(cls.viewer, "accounts.view_delegation")

        cls.stranger = User.objects.create_user("stranger.scope", "stranger.scope@example.test", PASSWORD)
        make_official(cls.stranger, delegation=cls.delegation, national_id="RUT-STRANGER")

    def test_viewer_can_list_but_not_create(self):
        self.client.force_login(self.viewer)

        self.assertEqual(self.client.get(reverse("delegation_list")).status_code, 200)

        response = self.client.post(
            reverse("delegation_create"), delegation_payload(delegation_name="Zona Prohibida")
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Delegation.objects.filter(delegation_name="Zona Prohibida").exists())

    def test_viewer_cannot_update_or_deactivate(self):
        self.client.force_login(self.viewer)

        self.assertEqual(
            self.client.get(reverse("delegation_update", args=[self.delegation.pk])).status_code, 403
        )
        self.assertEqual(
            self.client.post(reverse("delegation_deactivate", args=[self.delegation.pk])).status_code,
            403,
        )

        self.delegation.refresh_from_db()
        self.assertEqual(self.delegation.status, Delegation.Status.ACTIVE)

    def test_user_without_permission_is_rejected_and_audited(self):
        self.client.force_login(self.stranger)

        response = self.client.get(reverse("delegation_list"))

        self.assertEqual(response.status_code, 403)
        self.assertTrue(AuditLog.objects.filter(action="Acceso denegado").exists())


class DelegationRoleScopeTests(TestCase):
    """Un rol vigente de alcance ALL habilita la vista completa."""

    @classmethod
    def setUpTestData(cls):
        cls.delegation = Delegation.objects.create(
            delegation_name="Zona Rol", address="Calle Uno", commune="La Serena"
        )
        cls.role = Role.objects.create(role_name="Administrador de Prueba", scope=Role.Scope.ALL)
        cls.role.permissions.add(
            Permission.objects.get(content_type__app_label="accounts", codename="view_delegation")
        )

        cls.user = User.objects.create_user("rol.admin", "rol.admin@example.test", PASSWORD)
        official = make_official(
            cls.user, delegation=cls.delegation, national_id="RUT-ROL", email="rol.admin@example.test"
        )
        cls.role.officials.add(official)

    def test_role_all_scope_grants_full_visibility(self):
        Delegation.objects.create(
            delegation_name="Zona Otra", address="Calle Dos", commune="La Serena"
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("delegation_list"))

        self.assertContains(response, "Zona Rol")
        self.assertContains(response, "Zona Otra")
