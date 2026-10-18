from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COLUMN_WIDTH = 60


def _cell_value(value):
    """Valor seguro para una celda: sin caracteres de control y sin None."""
    if value is None:
        return ""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def build_xlsx(headers, rows, sheet_title="Registros"):
    """
    Libro Excel de una hoja: encabezados en negrita, fila de encabezado fija y autofiltro.

    No depende de ningún modelo: recibe los encabezados y las filas ya preparadas para
    mostrar (solo campos visibles para el usuario, nunca identificadores internos).
    Los textos se guardan siempre como texto, de modo que un valor que empiece con "="
    no se interpreta como fórmula al abrir el archivo.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]
    widths = [len(str(header)) for header in headers]

    sheet.append([_cell_value(header) for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row_number, row in enumerate(rows, start=2):
        for column_number, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_number, column=column_number, value=_cell_value(value))
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if column_number <= len(widths):
                widths[column_number - 1] = max(widths[column_number - 1], len(str(cell.value)))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, MAX_COLUMN_WIDTH)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def xlsx_response(filename, headers, rows, sheet_title="Registros"):
    """Respuesta de descarga con el libro generado por build_xlsx (`filename` sin extensión)."""
    response = HttpResponse(build_xlsx(headers, rows, sheet_title), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response
