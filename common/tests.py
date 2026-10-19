from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from common.exports import build_xlsx

User = get_user_model()

# Listados con exportación operativa (servidor) y maquetas sin persistencia.
SERVER_EXPORT_LISTS = (
    "delegation_list",
    "role_list",
    "user_maintainer",
    "official_maintainer",
    "item_maintainer",
    "attention_type_maintainer",
    "subattention_type_maintainer",
)
MOCKUP_LISTS = (
    "position_maintainer",
    "area_maintainer",
    "group_maintainer",
    "period_maintainer",
    "activity_maintainer",
)


class BuildXlsxTests(SimpleTestCase):
    def test_header_is_bold_frozen_and_filterable(self):
        sheet = load_workbook(BytesIO(build_xlsx(["Código", "Nombre"], [["A-1", "Uno"]]))).active

        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:B2")

    def test_text_is_never_read_as_a_formula_and_control_characters_are_removed(self):
        content = build_xlsx(["Nombre"], [["=1+1"], ["Línea\x07con control"], [None]])
        rows = list(load_workbook(BytesIO(content)).active.iter_rows(min_row=2, values_only=True))

        self.assertEqual(rows, [("=1+1",), ("Líneacon control",), (None,)])
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet["A2"].data_type, "s")


class ExportLabelTests(TestCase):
    """Todos los listados usan el mismo texto y ninguno ofrece un CSV generado en el navegador."""

    def setUp(self):
        self.client.force_login(User.objects.create_superuser("labels", "labels@example.test", "Clave-Labels-1"))

    def test_lists_offer_excel_export_and_never_csv(self):
        for route_name in SERVER_EXPORT_LISTS + MOCKUP_LISTS:
            with self.subTest(route=route_name):
                response = self.client.get(reverse(route_name))

                self.assertContains(response, "Exportar a Excel")
                self.assertNotContains(response, "CSV")
                self.assertNotContains(response, "Exportar Excel")

    def test_mockups_do_not_offer_an_operational_export(self):
        for route_name in MOCKUP_LISTS:
            with self.subTest(route=route_name):
                response = self.client.get(reverse(route_name))

                self.assertContains(
                    response, '<button class="btn maintainer-export-button" type="button" disabled>'
                )


class PaginationImportTests(SimpleTestCase):
    def test_old_import_paths_point_to_the_shared_mixin(self):
        from accounts.views import ClampedPaginationMixin as from_accounts
        from common.pagination import ClampedPaginationMixin
        from configuration.views import ClampedPaginationMixin as from_configuration

        self.assertIs(from_accounts, ClampedPaginationMixin)
        self.assertIs(from_configuration, ClampedPaginationMixin)
