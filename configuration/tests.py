from datetime import date
from io import BytesIO

from django.contrib.auth.models import Permission, User
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Delegation, Official, ServiceRecipient
from activities.models import SocialAttention
from configuration.admin import PeriodAdmin
from configuration.models import AttentionType, AuditLog, Period, SubattentionType


class ConfigurationFrontendRouteTests(TestCase):
    def setUp(self):
        self.client.force_login(
            User.objects.create_superuser("admin", "admin@example.test", "Clave-Admin-1")
        )

    def test_c4_and_c5_catalogue_routes_render_their_templates(self):
        routes_and_templates = (
            ("item_maintainer", "configuration/item_list.html", "Ítems y metas"),
            (
                "attention_type_maintainer",
                "configuration/attention_type_list.html",
                "Tipos de Atención",
            ),
            (
                "subattention_type_maintainer",
                "configuration/subattention_type_list.html",
                "Tipo de Atención asociado",
            ),
            ("period_maintainer", "configuration/period_list.html", "Períodos"),
        )

        for route_name, template_name, expected_content in routes_and_templates:
            with self.subTest(route=route_name):
                response = self.client.get(reverse(route_name))

                self.assertEqual(response.status_code, 200)
                self.assertTemplateUsed(response, template_name)
                self.assertContains(response, expected_content)
                self.assertContains(response, 'id="maintainer-table"')

    def test_c5_catalogues_show_persisted_crud_controls(self):
        response = self.client.get(reverse("attention_type_maintainer"))

        self.assertContains(response, f'href="{reverse("attention_type_create")}"')
        self.assertContains(response, "Exportar a Excel")


class AttentionCatalogCrudTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(
            "catalog.admin", "catalog.admin@example.test", "Clave-Admin-1"
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_attention_type_list_searches_filters_and_shows_model_data(self):
        active = AttentionType.objects.create(name="Orientación familiar", description="Apoyo")
        AttentionType.objects.create(name="Atención jurídica", is_active=False)

        response = self.client.get(
            reverse("attention_type_maintainer"), {"q": "familiar", "status": "1"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, active.code)
        self.assertContains(response, "Orientación familiar")
        self.assertNotContains(response, "Atención jurídica")

    def test_create_attention_type_generates_code_and_audit_entry(self):
        response = self.client.post(
            reverse("attention_type_create"),
            {"name": "Atención comunitaria", "description": "Apoyo local", "is_active": "True"},
            follow=True,
        )

        record = AttentionType.objects.get(name="Atención comunitaria")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(record.code, f"TAT-{record.pk:04d}")
        self.assertContains(response, "fue creado correctamente")
        self.assertTrue(AuditLog.objects.filter(action="Creación de tipo de atención").exists())

    def test_attention_type_duplicate_name_is_rejected_case_insensitively(self):
        AttentionType.objects.create(name="Apoyo social")

        response = self.client.post(
            reverse("attention_type_create"),
            {"name": " apoyo SOCIAL ", "description": "", "is_active": "True"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ya existe un tipo de atención con ese nombre.")
        self.assertEqual(AttentionType.objects.count(), 1)

    def test_update_attention_type_preserves_generated_code(self):
        record = AttentionType.objects.create(name="Orientación")
        original_code = record.code

        response = self.client.post(
            reverse("attention_type_update", args=(record.pk,)),
            {"name": "Orientación especializada", "description": "Actualizada", "is_active": "False"},
            follow=True,
        )

        record.refresh_from_db()
        self.assertEqual(record.name, "Orientación especializada")
        self.assertEqual(record.description, "Actualizada")
        self.assertFalse(record.is_active)
        self.assertEqual(record.code, original_code)
        self.assertContains(response, "fue actualizado correctamente")

    def test_subattention_creation_requires_an_active_attention_type(self):
        inactive_type = AttentionType.objects.create(name="Tipo inactivo", is_active=False)

        response = self.client.post(
            reverse("subattention_type_create"),
            {
                "name": "Subatención",
                "attention_type": inactive_type.pk,
                "is_active": "True",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "No puedes activar una subatención asociada a un tipo de atención inactivo.")
        self.assertFalse(SubattentionType.objects.exists())

    def test_subattention_create_saves_association_and_generated_code(self):
        attention_type = AttentionType.objects.create(name="Apoyo social")

        response = self.client.post(
            reverse("subattention_type_create"),
            {
                "name": "Orientación familiar",
                "attention_type": attention_type.pk,
                "is_active": "True",
            },
            follow=True,
        )

        subattention = SubattentionType.objects.get(name="Orientación familiar")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(subattention.attention_type, attention_type)
        self.assertEqual(subattention.code, f"SAT-{subattention.pk:04d}")
        self.assertContains(response, "fue creada correctamente")

    def test_attention_catalogue_detail_pages_render_related_data(self):
        attention_type = AttentionType.objects.create(name="Apoyo social")
        subattention = SubattentionType.objects.create(
            name="Orientación familiar", attention_type=attention_type
        )

        attention_response = self.client.get(
            reverse("attention_type_detail", args=(attention_type.pk,))
        )
        subattention_response = self.client.get(
            reverse("subattention_type_detail", args=(subattention.pk,))
        )

        self.assertEqual(attention_response.status_code, 200)
        self.assertTemplateUsed(attention_response, "configuration/attention_type_detail.html")
        self.assertContains(attention_response, attention_type.code)
        self.assertEqual(subattention_response.status_code, 200)
        self.assertTemplateUsed(
            subattention_response, "configuration/subattention_type_detail.html"
        )
        self.assertContains(subattention_response, attention_type.name)
        self.assertContains(subattention_response, subattention.code)

    def test_subattention_name_must_be_unique_within_attention_type(self):
        attention_type = AttentionType.objects.create(name="Apoyo social")
        SubattentionType.objects.create(name="Orientación", attention_type=attention_type)

        response = self.client.post(
            reverse("subattention_type_create"),
            {
                "name": "orientación",
                "attention_type": attention_type.pk,
                "is_active": "True",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ya existe una subatención con ese nombre para este tipo.")
        self.assertEqual(SubattentionType.objects.count(), 1)

    def test_attention_type_with_active_subattentions_cannot_be_deactivated(self):
        attention_type = AttentionType.objects.create(name="Apoyo social")
        SubattentionType.objects.create(name="Orientación", attention_type=attention_type)

        response = self.client.post(
            reverse("attention_type_deactivate", args=(attention_type.pk,)), follow=True
        )

        attention_type.refresh_from_db()
        self.assertTrue(attention_type.is_active)
        self.assertContains(response, "No se puede desactivar este tipo mientras tenga subatenciones activas.")

    def test_attention_type_cannot_be_deactivated_through_edit_with_active_subattentions(self):
        attention_type = AttentionType.objects.create(name="Apoyo social")
        SubattentionType.objects.create(name="Orientación", attention_type=attention_type)

        response = self.client.post(
            reverse("attention_type_update", args=(attention_type.pk,)),
            {"name": "Apoyo social", "description": "", "is_active": "False"},
        )

        attention_type.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(attention_type.is_active)
        self.assertContains(response, "No se puede desactivar este tipo mientras tenga subatenciones activas.")

    def test_subattention_type_association_cannot_change_after_historical_attention(self):
        original_type = AttentionType.objects.create(name="Apoyo social")
        other_type = AttentionType.objects.create(name="Orientación jurídica")
        subattention = SubattentionType.objects.create(
            name="Orientación familiar", attention_type=original_type
        )
        delegation = Delegation.objects.create(
            delegation_name="Centro", address="Calle 1", commune="La Serena"
        )
        official = Official.objects.create(
            full_name="Funcionario de prueba",
            national_id="RUT-CATALOG-HISTORY",
            email="catalog.history@example.test",
            delegation=delegation,
        )
        recipient = ServiceRecipient.objects.create(
            full_name="Persona de prueba", national_id="RUT-RECIP-HISTORY"
        )
        SocialAttention.objects.create(
            recipient=recipient,
            attention_type=original_type,
            subattention_type=subattention,
            delegation=delegation,
            responsible=official,
        )

        response = self.client.post(
            reverse("subattention_type_update", args=(subattention.pk,)),
            {
                "name": subattention.name,
                "attention_type": other_type.pk,
                "is_active": "True",
            },
        )

        subattention.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(subattention.attention_type, original_type)
        self.assertContains(
            response,
            "No puedes cambiar el tipo asociado porque esta subatención ya tiene atenciones registradas.",
        )

    def test_deactivation_is_logical_and_retains_audit_history(self):
        attention_type = AttentionType.objects.create(name="Apoyo social")
        subattention = SubattentionType.objects.create(name="Orientación", attention_type=attention_type)

        response = self.client.post(
            reverse("subattention_type_deactivate", args=(subattention.pk,)), follow=True
        )

        subattention.refresh_from_db()
        self.assertFalse(subattention.is_active)
        self.assertTrue(SubattentionType.objects.filter(pk=subattention.pk).exists())
        self.assertContains(response, "La subatención Orientación fue desactivada correctamente.")
        self.assertTrue(AuditLog.objects.filter(action="Desactivación de subatención").exists())

    def test_subattention_cannot_be_activated_while_parent_is_inactive(self):
        attention_type = AttentionType.objects.create(name="Apoyo social", is_active=False)
        subattention = SubattentionType.objects.create(
            name="Orientación", attention_type=attention_type, is_active=False
        )

        response = self.client.post(
            reverse("subattention_type_activate", args=(subattention.pk,)), follow=True
        )

        subattention.refresh_from_db()
        self.assertFalse(subattention.is_active)
        self.assertContains(response, "Activa primero el tipo de atención asociado.")

    def test_excel_export_respects_search_and_status_filters(self):
        AttentionType.objects.create(name="Atención activa", description="Comunitaria")
        AttentionType.objects.create(name="Atención inactiva", is_active=False)

        response = self.client.get(
            reverse("attention_type_export"), {"q": "activa", "status": "1"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        self.assertIn("Atención activa", str(rows))
        self.assertNotIn("Atención inactiva", str(rows))

    def test_subattention_excel_export_includes_associated_attention_type(self):
        attention_type = AttentionType.objects.create(name="Atención filtrada")
        subattention = SubattentionType.objects.create(
            name="Orientación exportada", attention_type=attention_type
        )

        response = self.client.get(
            reverse("subattention_type_export"), {"q": "Orientación exportada"}
        )

        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        self.assertEqual(
            rows[-1],
            (subattention.code, "Orientación exportada", "Atención filtrada", "Activo"),
        )

    def test_create_requires_add_permission(self):
        user = User.objects.create_user("catalog.reader", password="Clave-Reader-1")
        Official.objects.create(
            user=user,
            full_name="Lector de catálogos",
            national_id="RUT-CATALOG-READER",
            email="catalog.reader@example.test",
        )
        view_permission = Permission.objects.get(
            content_type__app_label="configuration", codename="view_attentiontype"
        )
        user.user_permissions.add(view_permission)
        self.client.force_login(user)

        response = self.client.get(reverse("attention_type_create"))

        self.assertEqual(response.status_code, 403)


class PeriodValidationTests(TestCase):
    """Validación Period.clean(): fechas coherentes y sin solapamientos."""

    def test_period_end_date_before_start_date_is_rejected(self):
        period = Period(start_date=date(2026, 3, 2), end_date=date(2026, 3, 1))

        with self.assertRaises(ValidationError) as error:
            period.full_clean()

        self.assertIn("end_date", error.exception.message_dict)

    def test_period_overlapping_existing_one_is_rejected(self):
        Period.objects.create(start_date=date(2026, 1, 1), end_date=date(2026, 1, 31))
        period = Period(start_date=date(2026, 1, 31), end_date=date(2026, 2, 28))

        with self.assertRaises(ValidationError) as error:
            period.full_clean()

        self.assertIn("start_date", error.exception.message_dict)

    def test_period_consecutive_to_existing_one_is_accepted(self):
        Period.objects.create(start_date=date(2026, 1, 1), end_date=date(2026, 1, 31))
        period = Period(start_date=date(2026, 2, 1), end_date=date(2026, 2, 28))

        period.full_clean()  


class PeriodActionTests(TestCase):
    """Acción del Admin 'Cerrar períodos seleccionados' y el método que usa."""

    def test_close_periods_action_is_registered_in_period_admin(self):
        self.assertIn("close_periods", PeriodAdmin.actions)

    def test_close_period_sets_status_to_closed(self):
        period = Period.objects.create(start_date=date(2026, 1, 1), end_date=date(2026, 1, 31))

        period.close_period()
        period.refresh_from_db()

        self.assertEqual(period.status, Period.Status.CLOSED)

class CatalogAuditAndExportCleanupTests(TestCase):
    """Auditoría por campo en Tipo y Sub Atención y exportación sin permiso por defecto."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("catalog.audit", "catalog.audit@example.test", "Clave-Admin-1")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_attention_type_update_audits_previous_and_new_value_per_field(self):
        record = AttentionType.objects.create(name="Orientación", description="Inicial")

        self.client.post(
            reverse("attention_type_update", args=(record.pk,)),
            {"name": "Orientación legal", "description": "Inicial", "is_active": "True"},
        )

        log = AuditLog.objects.get(action="Actualización de tipo de atención")
        self.assertEqual(log.previous_value, "Nombre: Orientación")
        self.assertEqual(log.new_value, "Nombre: Orientación legal")
        self.assertEqual(log.affected_record_id, record.pk)

    def test_subattention_update_audits_name_status_and_associated_type(self):
        first = AttentionType.objects.create(name="Social")
        second = AttentionType.objects.create(name="Jurídica")
        record = SubattentionType.objects.create(name="Orientación", attention_type=first)

        self.client.post(
            reverse("subattention_type_update", args=(record.pk,)),
            {"name": "Orientación inicial", "attention_type": second.pk, "is_active": "False"},
        )

        log = AuditLog.objects.get(action="Actualización de subatención")
        self.assertEqual(
            log.previous_value,
            f"Nombre: Orientación\nTipo de Atención asociado: {first.code} · Social\nEstado: Activo",
        )
        self.assertEqual(
            log.new_value,
            f"Nombre: Orientación inicial\nTipo de Atención asociado: {second.code} · Jurídica\nEstado: Inactivo",
        )

    def test_update_without_changes_is_audited_as_such(self):
        record = AttentionType.objects.create(name="Sin cambios", description="Igual")

        self.client.post(
            reverse("attention_type_update", args=(record.pk,)),
            {"name": "Sin cambios", "description": "Igual", "is_active": "True"},
        )

        log = AuditLog.objects.get(action="Actualización de tipo de atención")
        self.assertEqual((log.previous_value, log.new_value), ("Sin cambios", "Sin cambios"))

    def test_creation_audits_every_field(self):
        attention_type = AttentionType.objects.create(name="Social")

        self.client.post(
            reverse("subattention_type_create"),
            {"name": "Derivación", "attention_type": attention_type.pk, "is_active": "True"},
        )

        record = SubattentionType.objects.get(name="Derivación")
        log = AuditLog.objects.get(action="Creación de subatención")
        self.assertEqual(
            log.new_value,
            f"Código: {record.code}\nNombre: Derivación\nDescripción: —\n"
            f"Tipo de Atención asociado: {attention_type.code} · Social\nEstado: Activo",
        )

    def test_export_view_without_permission_required_fails_instead_of_staying_open(self):
        from django.core.exceptions import ImproperlyConfigured
        from django.test import RequestFactory

        from configuration.views import CatalogExportView

        class ForgottenPermissionExport(CatalogExportView):
            model = AttentionType
            filename = "olvido"
            headings = ("Código",)

            def row_values(self, record):
                return (record.code,)

        request = RequestFactory().get("/exportar/")
        request.user = self.admin

        with self.assertRaises(ImproperlyConfigured):
            ForgottenPermissionExport.as_view()(request)

    def test_exports_name_their_sheet_after_the_catalog(self):
        for route_name, sheet_title in (
            ("attention_type_export", "Tipos de atención"),
            ("subattention_type_export", "Subatenciones"),
        ):
            with self.subTest(route=route_name):
                response = self.client.get(reverse(route_name))

                self.assertEqual(load_workbook(BytesIO(response.content)).active.title, sheet_title)

    def test_out_of_range_page_is_clamped(self):
        for index in range(12):
            AttentionType.objects.create(name=f"Tipo {index:02d}")

        response = self.client.get(reverse("attention_type_maintainer"), {"page": 99})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_obj"].number, 2)
