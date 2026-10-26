"""
Pruebas del CRUD de Ítems.

Cubren creación (con código automático o manual), consulta, edición, desactivación y
reactivación lógicas, validaciones de código y nombre, permisos (403 con auditoría), búsqueda,
filtros, paginación, exportación a Excel y registro de auditoría con valor anterior y nuevo.
"""

from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import Official
from activities.models import Activity
from configuration.models import AuditLog, Item

User = get_user_model()
PASSWORD = "Clave-Segura-123"


def grant_permissions(user, *codenames):
    for codename in codenames:
        app_label, name = codename.split(".")
        permission = Permission.objects.get(content_type__app_label=app_label, codename=name)
        user.user_permissions.add(permission)
    return user


def make_official(user, *, status=Official.Status.ACTIVE, full_name=None, national_id=None, email=None):
    return Official.objects.create(
        user=user,
        full_name=full_name or user.username.replace(".", " ").title(),
        national_id=national_id or f"RUT-{user.username}",
        email=email or f"{user.username}@example.test",
        status=status,
    )


def item_payload(**overrides):
    payload = {
        "code": "",
        "item_name": "Ítem de Prueba",
        "description": "Ítem creado en las pruebas",
        "is_active": "True",
    }
    payload.update(overrides)
    return payload


class ItemCrudAdminTests(TestCase):
    """CRUD de ítems con un superusuario."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("admin.items", "admin.items@example.test", PASSWORD)
        make_official(cls.admin, full_name="Administrador Ítems", national_id="RUT-ADMIN-ITEMS")
        cls.item = Item.objects.create(item_name="Atención ciudadana", description="Atenciones en oficina")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_list_renders_persisted_items(self):
        response = self.client.get(reverse("item_maintainer"))

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "configuration/item_list.html")
        self.assertContains(response, "Atención ciudadana")
        self.assertContains(response, self.item.code)
        self.assertContains(response, f'href="{reverse("item_create")}"')

    def test_create_without_code_generates_it(self):
        response = self.client.post(reverse("item_create"), item_payload(item_name="Terreno"), follow=True)

        item = Item.objects.get(item_name="Terreno")
        self.assertEqual(item.code, f"ITM-{item.pk:04d}")
        self.assertTrue(item.is_active)
        self.assertContains(response, f"fue creado correctamente con el código {item.code}")
        log = AuditLog.objects.get(action="Creación de ítem", affected_record_id=item.pk)
        self.assertEqual(log.affected_table, Item._meta.db_table)
        self.assertIn(f"Código: {item.code}", log.new_value)

    def test_manual_code_is_uppercased_and_trimmed(self):
        self.client.post(
            reverse("item_create"),
            item_payload(code="  vis-01 ", item_name="  Visitas    domiciliarias  ", is_active="False"),
        )

        item = Item.objects.get(code="VIS-01")
        self.assertEqual(item.item_name, "Visitas domiciliarias")
        self.assertFalse(item.is_active)

    def test_duplicate_code_is_rejected_case_insensitively(self):
        Item.objects.create(item_name="Visitas", code="VIS-01")

        response = self.client.post(reverse("item_create"), item_payload(code="vis-01"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ya existe un ítem con este código")
        self.assertFalse(Item.objects.filter(item_name="Ítem de Prueba").exists())

    def test_reserved_automatic_code_format_is_rejected(self):
        response = self.client.post(reverse("item_create"), item_payload(code="itm-0099"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "se reservan para la numeración automática")
        self.assertFalse(Item.objects.filter(item_name="Ítem de Prueba").exists())

    def test_duplicate_name_is_rejected_case_insensitively(self):
        response = self.client.post(
            reverse("item_create"), item_payload(item_name="  ATENCIÓN   ciudadana ")
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ya existe un ítem con este nombre")
        self.assertEqual(Item.objects.count(), 1)

    def test_required_fields_are_validated(self):
        response = self.client.post(
            reverse("item_create"), {"code": "", "item_name": "   ", "description": "", "is_active": ""}
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ingresa el nombre del ítem")
        self.assertEqual(Item.objects.count(), 1)

    def test_update_persists_changes_and_audits_only_what_changed(self):
        original_code = self.item.code

        response = self.client.post(
            reverse("item_update", args=[self.item.pk]),
            item_payload(code=original_code, item_name="Atención presencial", description=self.item.description),
            follow=True,
        )

        self.item.refresh_from_db()
        self.assertEqual(self.item.item_name, "Atención presencial")
        self.assertEqual(self.item.code, original_code)
        self.assertContains(response, "fue actualizado correctamente")
        log = AuditLog.objects.get(action="Actualización de ítem", affected_record_id=self.item.pk)
        self.assertEqual(log.previous_value, "Nombre: Atención ciudadana")
        self.assertEqual(log.new_value, "Nombre: Atención presencial")

    def test_update_can_change_the_code_and_audits_it(self):
        previous_code = self.item.code

        self.client.post(
            reverse("item_update", args=[self.item.pk]),
            item_payload(code="atc-01", item_name=self.item.item_name, description=self.item.description),
        )

        self.item.refresh_from_db()
        self.assertEqual(self.item.code, "ATC-01")
        log = AuditLog.objects.get(action="Actualización de ítem", affected_record_id=self.item.pk)
        self.assertEqual(log.previous_value, f"Código: {previous_code}")
        self.assertEqual(log.new_value, "Código: ATC-01")

    def test_existing_code_cannot_be_cleared_on_edit(self):
        self.item.code = "MANUAL-1"
        self.item.save()

        for empty in ("", "   "):
            with self.subTest(code=repr(empty)):
                response = self.client.post(
                    reverse("item_update", args=[self.item.pk]),
                    item_payload(code=empty, item_name=self.item.item_name, description=self.item.description),
                )

                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "un ítem existente no puede quedar sin código")
                self.item.refresh_from_db()
                self.assertEqual(self.item.code, "MANUAL-1")
        self.assertFalse(AuditLog.objects.filter(action="Actualización de ítem").exists())

    def test_edit_form_marks_the_code_as_required(self):
        response = self.client.get(
            reverse("item_update", args=[self.item.pk]), HTTP_X_MODAL_REQUEST="true"
        )

        self.assertContains(response, f'value="{self.item.code}"')
        self.assertRegex(response.content.decode(), r'<input[^>]*name="code"[^>]*required')

    def test_update_form_does_not_change_status(self):
        self.client.post(
            reverse("item_update", args=[self.item.pk]),
            item_payload(code=self.item.code, item_name=self.item.item_name, is_active="False"),
        )

        self.item.refresh_from_db()
        self.assertTrue(self.item.is_active)

    def test_deactivate_is_logical_and_audited(self):
        response = self.client.post(reverse("item_deactivate", args=[self.item.pk]), follow=True)

        self.item.refresh_from_db()
        self.assertFalse(self.item.is_active)
        self.assertTrue(Item.objects.filter(pk=self.item.pk).exists())
        self.assertContains(response, "fue desactivado correctamente")
        log = AuditLog.objects.get(action="Desactivación de ítem", affected_record_id=self.item.pk)
        self.assertEqual(
            (log.previous_value, log.new_value),
            ("Estado: Activo", "Estado: Inactivo\nActividades registradas con el ítem: 0"),
        )

    def test_activate_after_deactivate(self):
        self.item.is_active = False
        self.item.save()

        response = self.client.post(reverse("item_activate", args=[self.item.pk]), follow=True)

        self.item.refresh_from_db()
        self.assertTrue(self.item.is_active)
        self.assertContains(response, "fue activado correctamente")
        self.assertTrue(AuditLog.objects.filter(action="Activación de ítem", affected_record_id=self.item.pk).exists())

    def test_deactivate_and_activate_reject_get_requests(self):
        for route_name in ("item_deactivate", "item_activate"):
            with self.subTest(route=route_name):
                response = self.client.get(reverse(route_name, args=[self.item.pk]))
                self.assertEqual(response.status_code, 405)

        self.item.refresh_from_db()
        self.assertTrue(self.item.is_active)

    def test_list_offers_confirmation_before_deactivating(self):
        response = self.client.get(reverse("item_maintainer"))

        self.assertContains(response, 'data-swal-title="Desactivar ítem"')
        self.assertContains(response, f'action="{reverse("item_deactivate", args=[self.item.pk])}"')

    def test_search_is_partial_and_case_insensitive_across_fields(self):
        Item.objects.create(item_name="Reunión de coordinación", code="REU-01")

        cases = (
            ("CIUDADANA", "Atención ciudadana", "Reunión de coordinación"),
            ("en oficina", "Atención ciudadana", "Reunión de coordinación"),
            ("reu-0", "Reunión de coordinación", "Atención ciudadana"),
        )
        for query, expected, excluded in cases:
            with self.subTest(query=query):
                response = self.client.get(reverse("item_maintainer"), {"q": query})
                self.assertContains(response, expected)
                self.assertNotContains(response, excluded)

    def test_status_filter_combines_with_search(self):
        Item.objects.create(item_name="Atención telefónica", is_active=False)

        response = self.client.get(reverse("item_maintainer"), {"q": "atención", "status": "INACTIVE"})

        self.assertContains(response, "Atención telefónica")
        self.assertNotContains(response, "Atención ciudadana")

    def test_empty_states_distinguish_no_matches(self):
        response = self.client.get(reverse("item_maintainer"), {"q": "no-existe"})
        self.assertContains(response, "No hay ítems que coincidan con la búsqueda o el filtro aplicado.")

        Item.objects.all().delete()
        response = self.client.get(reverse("item_maintainer"))
        self.assertContains(response, "Aún no hay ítems registrados.")

    def test_pagination_is_ten_per_page_and_keeps_filters(self):
        for index in range(12):
            Item.objects.create(item_name=f"Ítem Paginado {index:02d}")

        first_page = self.client.get(reverse("item_maintainer"), {"q": "paginado"})

        self.assertEqual(len(first_page.context["items"]), 10)
        self.assertEqual(first_page.context["paginator"].count, 12)
        self.assertEqual(first_page.context["querystring"], "q=paginado")
        self.assertContains(first_page, "q=paginado&amp;page=2")
        self.assertContains(first_page, "de 12 registros")

        second_page = self.client.get(reverse("item_maintainer"), {"q": "paginado", "page": 2})

        self.assertEqual(len(second_page.context["items"]), 2)

    def test_out_of_range_page_is_clamped(self):
        response = self.client.get(reverse("item_maintainer"), {"page": 999})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_obj"].number, 1)

    def test_list_and_export_are_ordered_by_code(self):
        Item.objects.all().delete()  # su código automático podría coincidir con los de la prueba
        # El orden por nombre (Atención, Bitácora, Control, Reunión, Terreno) difiere del de código.
        for code, name in (
            ("ITM-0003", "Atención"),
            ("ITM-0001", "Terreno"),
            ("VIS-01", "Bitácora"),
            ("ITM-0010", "Control"),
            ("ITM-0002", "Reunión"),
        ):
            Item.objects.create(code=code, item_name=name)
        expected = ["ITM-0001", "ITM-0002", "ITM-0003", "ITM-0010", "VIS-01"]

        response = self.client.get(reverse("item_maintainer"))
        self.assertEqual([item.code for item in response.context["items"]], expected)

        export = self.client.get(reverse("item_export"))
        rows = list(load_workbook(BytesIO(export.content)).active.iter_rows(values_only=True))
        self.assertEqual([row[0] for row in rows[1:]], expected)

    def test_export_is_xlsx_with_filtered_rows_and_no_internal_fields(self):
        for index in range(12):
            Item.objects.create(item_name=f"Ítem Exportado {index:02d}")
        Item.objects.create(item_name="Ítem Exportado Inactivo", is_active=False)

        response = self.client.get(reverse("item_export"), {"q": "exportado", "status": "ACTIVE"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('attachment; filename="items.xlsx"', response["Content-Disposition"])
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Código", "Nombre", "Descripción", "Estado"))
        self.assertEqual(len(rows), 13)
        self.assertNotIn("Ítem Exportado Inactivo", [row[1] for row in rows])
        exported = Item.objects.get(item_name="Ítem Exportado 00")
        self.assertEqual((rows[1][0], rows[1][1], rows[1][3]), (exported.code, "Ítem Exportado 00", "Activo"))

    def test_create_update_and_detail_views_render(self):
        create_response = self.client.get(reverse("item_create"), HTTP_X_MODAL_REQUEST="true")
        self.assertEqual(create_response.status_code, 200)
        self.assertTemplateUsed(create_response, "configuration/item_form.html")

        update_response = self.client.get(
            reverse("item_update", args=[self.item.pk]), HTTP_X_MODAL_REQUEST="true"
        )
        self.assertEqual(update_response.status_code, 200)
        self.assertContains(update_response, f"Editar ítem {self.item.code}")

        detail_response = self.client.get(reverse("item_detail", args=[self.item.pk]))
        self.assertEqual(detail_response.status_code, 200)
        self.assertTemplateUsed(detail_response, "configuration/item_detail.html")
        self.assertContains(detail_response, "Atenciones en oficina")

    def test_anonymous_is_redirected_to_login(self):
        self.client.logout()
        url = reverse("item_maintainer")

        self.assertRedirects(self.client.get(url), f"{reverse('login')}?next={url}")


class ItemPermissionTests(TestCase):
    """Cada operación exige su permiso; los rechazos responden 403 y quedan auditados."""

    @classmethod
    def setUpTestData(cls):
        cls.item = Item.objects.create(item_name="Ítem Existente")

        cls.viewer = User.objects.create_user("viewer.items", "viewer.items@example.test", PASSWORD)
        make_official(cls.viewer)
        grant_permissions(cls.viewer, "configuration.view_item")

        cls.creator = User.objects.create_user("creator.items", "creator.items@example.test", PASSWORD)
        make_official(cls.creator)
        grant_permissions(cls.creator, "configuration.view_item", "configuration.add_item")

        cls.stranger = User.objects.create_user("stranger.items", "stranger.items@example.test", PASSWORD)
        make_official(cls.stranger)

    def test_viewer_can_list_export_and_view_but_not_modify(self):
        self.client.force_login(self.viewer)

        self.assertEqual(self.client.get(reverse("item_maintainer")).status_code, 200)
        self.assertEqual(self.client.get(reverse("item_export")).status_code, 200)
        self.assertEqual(self.client.get(reverse("item_detail", args=[self.item.pk])).status_code, 200)
        self.assertEqual(self.client.post(reverse("item_create"), item_payload()).status_code, 403)
        self.assertEqual(self.client.get(reverse("item_update", args=[self.item.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse("item_deactivate", args=[self.item.pk])).status_code, 403)

        response = self.client.get(reverse("item_maintainer"))
        self.assertNotContains(response, reverse("item_create"))
        self.assertNotContains(response, reverse("item_update", args=[self.item.pk]))
        self.item.refresh_from_db()
        self.assertTrue(self.item.is_active)
        self.assertFalse(Item.objects.filter(item_name="Ítem de Prueba").exists())
        self.assertEqual(AuditLog.objects.filter(action="Acceso denegado").count(), 3)

    def test_creator_can_create_but_not_edit(self):
        self.client.force_login(self.creator)

        self.client.post(reverse("item_create"), item_payload())

        self.assertTrue(Item.objects.filter(item_name="Ítem de Prueba").exists())
        self.assertEqual(self.client.get(reverse("item_update", args=[self.item.pk])).status_code, 403)

    def test_user_without_permission_is_rejected_and_audited(self):
        self.client.force_login(self.stranger)

        self.assertEqual(self.client.get(reverse("item_maintainer")).status_code, 403)
        self.assertEqual(self.client.get(reverse("item_export")).status_code, 403)
        log = AuditLog.objects.filter(action="Acceso denegado").first()
        self.assertEqual(log.official, self.stranger.official)


class ItemDeactivationWithActivitiesTests(TestCase):
    """Como Roles con sus funcionarios: desactivar informa cuántas actividades usan el ítem."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser("admin.uso", "admin.uso@example.test", PASSWORD)
        official = make_official(cls.admin, national_id="RUT-ADMIN-USO")
        cls.item = Item.objects.create(item_name="Visitas a terreno")
        for day in (1, 2):
            Activity.objects.create(
                official=official, item=cls.item, date=date(2026, 10, day), record_type="Terreno"
            )
        Activity.objects.create(
            official=official,
            item=cls.item,
            date=date(2026, 10, 3),
            record_type="Terreno",
            deleted_at=timezone.now(),
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_confirmation_shows_how_many_activities_use_the_item(self):
        response = self.client.get(reverse("item_maintainer"))

        self.assertContains(response, "2 actividades registradas lo usan")
        self.assertContains(response, 'name="affected_activities" value="2"')

        detail = self.client.get(reverse("item_detail", args=[self.item.pk]))
        self.assertContains(detail, "2 actividades registradas lo usan")

    def test_outdated_count_does_not_deactivate(self):
        response = self.client.post(
            reverse("item_deactivate", args=[self.item.pk]), {"affected_activities": "1"}, follow=True
        )

        self.item.refresh_from_db()
        self.assertTrue(self.item.is_active)
        self.assertContains(response, "tiene 2 actividades registradas")

    def test_confirmed_count_deactivates_and_audits_it(self):
        response = self.client.post(
            reverse("item_deactivate", args=[self.item.pk]), {"affected_activities": "2"}, follow=True
        )

        self.item.refresh_from_db()
        self.assertFalse(self.item.is_active)
        self.assertEqual(self.item.activities.count(), 3)
        self.assertContains(response, "conservan el ítem en su historial")
        log = AuditLog.objects.get(action="Desactivación de ítem")
        self.assertIn("Actividades registradas con el ítem: 2", log.new_value)
